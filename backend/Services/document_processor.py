from __future__ import annotations

import hashlib
import html
import logging
import math
import os
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from pptx import Presentation
except Exception:  # pragma: no cover
    Presentation = None

try:  # Optional scanned-PDF OCR stack.
    import fitz
    import pytesseract
    from PIL import Image
except Exception:  # pragma: no cover
    fitz = None
    pytesseract = None
    Image = None

logger = logging.getLogger(__name__)

SUPPORTED_FORMATS = {
    "pdf": [".pdf"],
    "word": [".docx"],
    "powerpoint": [".pptx"],
    "text": [".txt", ".md"],
    "html": [".html", ".htm"],
}

CHUNK_SIZE = 1400
CHUNK_OVERLAP = 180


def _positive_int_setting(name: str, default: int) -> int:
    try:
        value = int(os.getenv(name, str(default)))
        return value if value > 0 else default
    except (TypeError, ValueError):
        return default


MAX_EXTRACTED_CHARACTERS = _positive_int_setting("DOCUMENT_MAX_EXTRACTED_CHARACTERS", 5_000_000)
MAX_OOXML_EXPANDED_BYTES = _positive_int_setting("DOCUMENT_MAX_OOXML_EXPANDED_BYTES", 250 * 1024 * 1024)
MAX_OOXML_ENTRIES = _positive_int_setting("DOCUMENT_MAX_OOXML_ENTRIES", 10_000)
MAX_PDF_PAGES = _positive_int_setting("DOCUMENT_MAX_PDF_PAGES", 2_000)
MAX_OCR_PAGES = _positive_int_setting("DOCUMENT_MAX_OCR_PAGES", 100)
EMBEDDING_BATCH_SIZE = _positive_int_setting("DOCUMENT_EMBEDDING_BATCH_SIZE", 64)
NEO4J_WRITE_BATCH_SIZE = _positive_int_setting("DOCUMENT_WRITE_BATCH_SIZE", 100)
EMBEDDING_VECTOR_DIMENSIONS = _positive_int_setting("EMBEDDING_VECTOR_DIMENSIONS", 768)
DATASHEET_INDEX_NAME = os.getenv("NEO4J_DATASHEET_VECTOR_INDEX", "datasheet_index").strip() or "datasheet_index"
DATASHEET_KEYWORD_INDEX_NAME = os.getenv("NEO4J_DATASHEET_KEYWORD_INDEX", "datasheetkeyword").strip() or "datasheetkeyword"


def _neo4j_identifier(value: str, setting_name: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,127}", value):
        raise RuntimeError(f"{setting_name} must be a valid Neo4j identifier")
    return value


DATASHEET_INDEX_NAME = _neo4j_identifier(DATASHEET_INDEX_NAME, "NEO4J_DATASHEET_VECTOR_INDEX")
DATASHEET_KEYWORD_INDEX_NAME = _neo4j_identifier(
    DATASHEET_KEYWORD_INDEX_NAME,
    "NEO4J_DATASHEET_KEYWORD_INDEX",
)
UPSERT_DATASHEET_CHUNK = """
UNWIND $rows AS row
MERGE (d:DocumentAsset {document_id: row.document_id})
SET d.filename = row.filename,
    d.source = row.source,
    d.file_type = row.file_type,
    d.updated_at = datetime(),
    d.content_hash = row.content_hash,
    d.chunk_count = row.chunk_count,
    d.status = 'indexed'
MERGE (c:DatasheetChunk {document_id: row.document_id, chunk_id: row.chunk_id})
SET c.content = row.content,
    c.source = row.source,
    c.embedding = row.embedding,
    c.filename = row.filename,
    c.file_type = row.file_type,
    c.content_hash = row.content_hash,
    c.chunk_index = row.chunk_index,
    c.char_start = row.char_start,
    c.char_end = row.char_end,
    c.updated_at = datetime()
MERGE (d)-[:HAS_CHUNK]->(c)
"""

DOCUMENT_ASSET_CONSTRAINT = (
    "CREATE CONSTRAINT document_asset_id IF NOT EXISTS "
    "FOR (d:DocumentAsset) REQUIRE d.document_id IS UNIQUE"
)
DATASHEET_CHUNK_CONSTRAINT = (
    "CREATE CONSTRAINT datasheet_chunk_identity IF NOT EXISTS "
    "FOR (c:DatasheetChunk) REQUIRE (c.document_id, c.chunk_id) IS UNIQUE"
)


class DocumentProcessingCancelled(RuntimeError):
    pass


def get_format_description() -> dict[str, str]:
    return {
        "pdf": "Portable Document Format text extraction and DatasheetChunk indexing",
        "word": "Microsoft Word text extraction and DatasheetChunk indexing",
        "powerpoint": "Microsoft PowerPoint slide text extraction and DatasheetChunk indexing",
        "text": "Plain text or Markdown extraction and DatasheetChunk indexing",
        "html": "HTML text extraction and DatasheetChunk indexing",
    }


def cleanup_temp_directory() -> None:
    return None


def runtime_status() -> dict[str, Any]:
    errors: list[str] = []
    embedder_available = False
    try:
        from backend.core.llm import EMBEDDER_AVAILABLE as _available
        embedder_available = bool(_available)
    except Exception as exc:
        errors.append(f"llm import: {exc}")

    try:
        from backend.Services.graph_embeddings import ensure_indexes_standalone as _ensure_indexes  # noqa: F401
    except Exception as exc:
        errors.append(f"graph_embeddings import: {exc}")

    ocr_available = False
    if fitz is not None and pytesseract is not None and Image is not None:
        try:
            pytesseract.get_tesseract_version()
            ocr_available = True
        except Exception:
            ocr_available = False

    return {
        "available": embedder_available and not errors,
        "embedder_available": embedder_available,
        "ocr_available": ocr_available,
        "errors": errors,
    }


def _detect_format(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    for fmt, suffixes in SUPPORTED_FORMATS.items():
        if suffix in suffixes:
            return fmt
    raise ValueError(f"Unsupported document format: {suffix}")


def _extract_pdf_text(file_path: Path) -> str:
    if PdfReader is None:
        raise RuntimeError("pypdf is not available")
    reader = PdfReader(str(file_path))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDF exceeds the {MAX_PDF_PAGES} page extraction limit")
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            logger.warning("PDF page extraction failed for %s page %s: %s", file_path.name, index, exc)
            text = ""
        text = text.strip()
        if text:
            pages.append(f"[Page {index}]\n{text}")
    native_text = "\n\n".join(pages).strip()
    return native_text or _extract_pdf_ocr_text(file_path)


def _extract_pdf_ocr_text(file_path: Path) -> str:
    if fitz is None or pytesseract is None or Image is None:
        return ""
    document = fitz.open(str(file_path))
    try:
        if document.page_count > MAX_OCR_PAGES:
            raise ValueError(f"Scanned PDF exceeds the {MAX_OCR_PAGES} page OCR limit")
        pages: list[str] = []
        for index in range(document.page_count):
            page = document.load_page(index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            try:
                text = str(pytesseract.image_to_string(image) or "").strip()
            except Exception as exc:
                logger.warning("OCR unavailable for %s page %s: %s", file_path.name, index + 1, exc)
                return ""
            if text:
                pages.append(f"[Page {index + 1} OCR]\n{text}")
        return "\n\n".join(pages).strip()
    finally:
        document.close()


def _validate_ooxml_archive(file_path: Path, required_prefix: str) -> None:
    try:
        with zipfile.ZipFile(file_path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_OOXML_ENTRIES:
                raise ValueError(f"Office document exceeds the {MAX_OOXML_ENTRIES} entry limit")
            expanded_size = sum(max(0, entry.file_size) for entry in entries)
            if expanded_size > MAX_OOXML_EXPANDED_BYTES:
                raise ValueError("Office document exceeds the expanded-content safety limit")
            names = {entry.filename.replace("\\", "/") for entry in entries}
            if "[Content_Types].xml" not in names or not any(name.startswith(required_prefix) for name in names):
                raise ValueError(f"File is not a valid {file_path.suffix.lower()} Office document")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError("Encrypted Office documents are not supported")
    except zipfile.BadZipFile as exc:
        raise ValueError(f"File is not a valid Office archive: {file_path.name}") from exc


def _extract_docx_text(file_path: Path) -> str:
    if DocxDocument is None:
        raise RuntimeError("python-docx is not available")
    _validate_ooxml_archive(file_path, "word/")
    document = DocxDocument(str(file_path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    return "\n".join(paragraphs).strip()


def _extract_pptx_text(file_path: Path) -> str:
    if Presentation is None:
        raise RuntimeError("python-pptx is not available")
    _validate_ooxml_archive(file_path, "ppt/")
    presentation = Presentation(str(file_path))
    slides = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            text = text.strip()
            if text:
                parts.append(text)
        if parts:
            slides.append(f"[Slide {slide_index}]\n" + "\n".join(parts))
    return "\n\n".join(slides).strip()


def _extract_plain_text(file_path: Path) -> str:
    raw = file_path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unsupported text encoding: {file_path.name}")


class _VisibleHtmlTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"script", "style", "noscript", "template"}:
            self._ignored_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "template"} and self._ignored_depth:
            self._ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self.parts.append(data)


def _extract_html_text(file_path: Path) -> str:
    raw = _extract_plain_text(file_path)
    parser = _VisibleHtmlTextParser()
    parser.feed(raw)
    parser.close()
    return re.sub(r"\s+", " ", html.unescape(" ".join(parser.parts))).strip()


def extract_text_from_file(file_path: str | Path) -> dict[str, Any]:
    path = Path(file_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"Document not found: {path}")
    file_type = _detect_format(path)
    if file_type == "pdf":
        text = _extract_pdf_text(path)
    elif file_type == "word":
        text = _extract_docx_text(path)
    elif file_type == "powerpoint":
        text = _extract_pptx_text(path)
    elif file_type == "text":
        text = _extract_plain_text(path)
    elif file_type == "html":
        text = _extract_html_text(path)
    else:  # pragma: no cover
        raise ValueError(f"Unsupported document format: {path.suffix}")

    text = (text or "").strip()
    if not text:
        raise ValueError(f"No extractable text found in document: {path.name}")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError(
            f"Extracted text exceeds the {MAX_EXTRACTED_CHARACTERS} character processing limit"
        )

    return {
        "path": str(path),
        "filename": path.name,
        "file_type": file_type,
        "text": text,
    }


def _chunk_document_records(text: str) -> list[dict[str, Any]]:
    normalized = (text or "").strip()
    if not normalized:
        return []

    chunks: list[dict[str, Any]] = []
    start = 0
    while start < len(normalized):
        end = min(start + CHUNK_SIZE, len(normalized))
        if end < len(normalized):
            soft_break = normalized.rfind("\n", start, end)
            if soft_break <= start:
                soft_break = normalized.rfind(" ", start, end)
            if soft_break > start + 200:
                end = soft_break
        raw_chunk = normalized[start:end]
        leading = len(raw_chunk) - len(raw_chunk.lstrip())
        trailing = len(raw_chunk) - len(raw_chunk.rstrip())
        chunk = raw_chunk.strip()
        if chunk:
            chunks.append({
                "content": chunk,
                "char_start": start + leading,
                "char_end": end - trailing,
            })
        if end >= len(normalized):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def _chunk_document(text: str) -> list[str]:
    return [record["content"] for record in _chunk_document_records(text)]


def _hash_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _embed_texts(chunks: Iterable[str], cancel_check=None) -> list[list[float]]:
    normalized_chunks = list(chunks)
    from backend.core.llm import EMBEDDER_AVAILABLE, embeddings
    if not EMBEDDER_AVAILABLE:
        raise RuntimeError("Embeddings model is unavailable")
    vectors: list[list[float]] = []
    for start in range(0, len(normalized_chunks), EMBEDDING_BATCH_SIZE):
        if cancel_check and cancel_check():
            raise DocumentProcessingCancelled("Document processing was cancelled")
        vectors.extend(embeddings.embed_documents(normalized_chunks[start:start + EMBEDDING_BATCH_SIZE]))
    return vectors


_REQUIREMENT_ID_PATTERN = re.compile(r"\b(?:REQ|RQM|REQUIREMENT)[-_ ]?\d+[A-Za-z0-9_.-]*\b", re.IGNORECASE)
_PARAMETER_PATTERN = re.compile(
    r"\b(?P<value>[+-]?(?:\d+(?:\.\d+)?|\.\d+))\s*(?P<unit>mm|cm|m|km|kg|g|N|kN|Pa|kPa|MPa|V|A|Hz|rpm|°C|degC|s|ms|%)\b",
    re.IGNORECASE,
)


def extract_semantic_proposals(text: str, document_id: str, limit: int = 200) -> dict[str, Any]:
    """Extract deterministic, review-only requirement and parameter candidates."""
    normalized = str(text or "")
    entities: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    sentences = [item.strip() for item in re.split(r"(?<=[.!?])\s+|[\r\n]+", normalized) if item.strip()]
    for sentence in sentences:
        requirement_ids = _REQUIREMENT_ID_PATTERN.findall(sentence)
        requirement_language = bool(re.search(r"\b(shall|must|should|required to)\b", sentence, re.IGNORECASE))
        if requirement_ids or requirement_language:
            identifier = requirement_ids[0].upper().replace(" ", "-") if requirement_ids else _hash_text(sentence)[:16]
            key = ("Requirement", identifier)
            if key not in seen:
                seen.add(key)
                entities.append({
                    "candidate_id": f"{document_id}:requirement:{identifier}",
                    "entity_type": "Requirement",
                    "identifier": identifier,
                    "label": sentence[:240],
                    "source_text": sentence[:1000],
                    "confidence": 0.95 if requirement_ids else 0.78,
                    "mapping_proposal": {
                        "target_type": "oslc_rm:Requirement",
                        "relationship": "rdf:type",
                        "status": "proposed",
                    },
                })
        for match in _PARAMETER_PATTERN.finditer(sentence):
            token = f"{match.group('value')} {match.group('unit')}"
            key = ("Parameter", token.casefold())
            if key in seen:
                continue
            seen.add(key)
            entities.append({
                "candidate_id": f"{document_id}:parameter:{_hash_text(token.casefold())[:16]}",
                "entity_type": "Parameter",
                "label": token,
                "value": float(match.group("value")),
                "unit": match.group("unit"),
                "source_text": sentence[:1000],
                "confidence": 0.9,
                "mapping_proposal": {
                    "target_type": "Parameter",
                    "relationship": "rdf:type",
                    "status": "proposed",
                },
            })
        if len(entities) >= limit:
            break
    return {
        "document_id": document_id,
        "entities": entities[:limit],
        "entity_count": min(len(entities), limit),
        "truncated": len(entities) > limit,
        "requires_human_approval": True,
        "committed": False,
    }


def _upsert_datasheet_chunks(rows: list[dict[str, Any]]) -> None:
    from backend.core.db_config import get_config, get_driver
    config = get_config()
    driver = get_driver()
    with driver.session(database=config.database) as session:
        session.run(
            f"""
            CREATE VECTOR INDEX `{DATASHEET_INDEX_NAME}` IF NOT EXISTS
            FOR (c:DatasheetChunk) ON (c.embedding)
            OPTIONS {{indexConfig: {{
                `vector.dimensions`: {EMBEDDING_VECTOR_DIMENSIONS},
                `vector.similarity_function`: 'cosine'
            }}}}
            """
        ).consume()
        session.run(
            f"""
            CREATE FULLTEXT INDEX `{DATASHEET_KEYWORD_INDEX_NAME}` IF NOT EXISTS
            FOR (c:DatasheetChunk) ON EACH [c.content]
            """
        ).consume()
        index_record = session.run(
            "SHOW VECTOR INDEXES YIELD name, options WHERE name = $name RETURN options",
            name=DATASHEET_INDEX_NAME,
        ).single()
        options = dict(index_record["options"] or {}) if index_record else {}
        index_config = dict(options.get("indexConfig") or {})
        actual_dimension = index_config.get("vector.dimensions")
        if actual_dimension is not None and int(actual_dimension) != EMBEDDING_VECTOR_DIMENSIONS:
            raise RuntimeError(
                f"Vector index '{DATASHEET_INDEX_NAME}' dimension {actual_dimension} does not match "
                f"configured embedding dimension {EMBEDDING_VECTOR_DIMENSIONS}"
            )
        session.run(DOCUMENT_ASSET_CONSTRAINT).consume()
        session.run(DATASHEET_CHUNK_CONSTRAINT).consume()

        def write_chunks(transaction) -> None:
            for start in range(0, len(rows), NEO4J_WRITE_BATCH_SIZE):
                transaction.run(
                    UPSERT_DATASHEET_CHUNK,
                    rows=rows[start:start + NEO4J_WRITE_BATCH_SIZE],
                ).consume()

        session.execute_write(write_chunks)


def _validated_embeddings(chunks: list[str], embeddings: list[list[float]]) -> list[list[float]]:
    if len(embeddings) != len(chunks):
        raise RuntimeError(
            f"Embedding backend returned {len(embeddings)} vectors for {len(chunks)} chunks"
        )
    dimension = 0
    normalized: list[list[float]] = []
    for vector in embeddings:
        values = [float(value) for value in vector]
        if not values or any(not math.isfinite(value) for value in values):
            raise RuntimeError("Embedding backend returned an empty or non-finite vector")
        if dimension and len(values) != dimension:
            raise RuntimeError("Embedding backend returned inconsistent vector dimensions")
        dimension = dimension or len(values)
        normalized.append(values)
    if dimension and dimension != EMBEDDING_VECTOR_DIMENSIONS:
        raise RuntimeError(
            f"Embedding dimension {dimension} does not match configured vector index dimension "
            f"{EMBEDDING_VECTOR_DIMENSIONS}"
        )
    return normalized


def process_documents_batch(
    file_paths: list[str],
    index_name: str | None = None,
    *,
    cancel_check=None,
    progress_callback=None,
    include_semantic_proposals: bool = True,
) -> dict[str, Any]:
    normalized_index_name = str(index_name or DATASHEET_INDEX_NAME).strip()
    if normalized_index_name != DATASHEET_INDEX_NAME:
        raise ValueError(
            f"Unsupported document index '{normalized_index_name}'; configured index is '{DATASHEET_INDEX_NAME}'"
        )
    status = runtime_status()
    if not status["embedder_available"]:
        raise RuntimeError("Document processing requires a working embedding backend")

    results: list[dict[str, Any]] = []
    success_count = 0
    failed_count = 0

    for raw_path in file_paths:
        try:
            if cancel_check and cancel_check():
                raise DocumentProcessingCancelled("Document processing was cancelled")
            if progress_callback:
                progress_callback("extract", {"file": Path(raw_path).name})
            extracted = extract_text_from_file(raw_path)
            chunk_records = _chunk_document_records(extracted["text"])
            if not chunk_records:
                raise ValueError(f"No chunks could be created for document: {extracted['filename']}")
            chunks = [record["content"] for record in chunk_records]
            if progress_callback:
                progress_callback("embed", {"file": extracted["filename"], "chunks": len(chunks)})
            chunk_embeddings = _validated_embeddings(chunks, _embed_texts(chunks, cancel_check=cancel_check))
            content_hash = _hash_text(extracted["text"])
            document_id = _hash_text(f"{extracted['filename'].casefold()}\0{content_hash}")
            source_uri = f"document://{document_id}"
            semantic_proposals = (
                extract_semantic_proposals(extracted["text"], document_id)
                if include_semantic_proposals
                else {"entities": [], "entity_count": 0, "requires_human_approval": True, "committed": False}
            )
            rows = []
            for idx, (chunk_record, embedding) in enumerate(zip(chunk_records, chunk_embeddings), start=1):
                rows.append({
                    "document_id": document_id,
                    "source": source_uri,
                    "filename": extracted["filename"],
                    "file_type": extracted["file_type"],
                    "chunk_id": f"chunk_{idx:04d}",
                    "chunk_index": idx,
                    "chunk_count": len(chunks),
                    "content": chunk_record["content"],
                    "char_start": chunk_record["char_start"],
                    "char_end": chunk_record["char_end"],
                    "embedding": embedding,
                    "content_hash": content_hash,
                })
            if cancel_check and cancel_check():
                raise DocumentProcessingCancelled("Document processing was cancelled")
            if progress_callback:
                progress_callback("index", {"file": extracted["filename"], "chunks": len(rows)})
            _upsert_datasheet_chunks(rows)
            results.append({
                "file": extracted["filename"],
                "document_id": document_id,
                "source": source_uri,
                "file_type": extracted["file_type"],
                "status": "success",
                "chunks_created": len(rows),
                "index_name": normalized_index_name,
                "content_hash": content_hash,
                "semantic_proposals": semantic_proposals,
            })
            success_count += 1
        except DocumentProcessingCancelled:
            raise
        except Exception as exc:
            logger.exception("Document processing failed for %s", raw_path)
            error_message = str(exc)
            try:
                error_message = error_message.replace(str(Path(raw_path).resolve()), Path(raw_path).name)
            except Exception:
                pass
            results.append({
                "file": Path(raw_path).name,
                "status": "failed",
                "error": error_message,
                "index_name": normalized_index_name,
            })
            failed_count += 1
        finally:
            if progress_callback:
                progress_callback("file_complete", {"file": Path(raw_path).name})

    return {
        "summary": {
            "total_files": len(file_paths),
            "successfully_processed": success_count,
            "failed_processing": failed_count,
        },
        "processing_results": results,
    }
