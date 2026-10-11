"""Build the Word environment reference from its reviewed Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, RGBColor

root = Path(__file__).resolve().parents[1]
source = root / 'deliverables/DEPO_Environment_Configuration_Reference.md'
target = source.with_suffix('.docx')
doc = Document(target)
body = doc.element.body
for child in list(body):
    if not child.tag.endswith('sectPr'):
        body.remove(child)
for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(style.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
code = False
for line in source.read_text(encoding='utf-8').splitlines():
    if line.startswith('```'):
        code = not code
        continue
    if not line:
        continue
    match = re.match(r'^(#{1,4}) (.*)$', line)
    if match and not code:
        level, text = len(match[1]), match[2]
        p = doc.add_paragraph(text, style='Title' if level == 1 else f'Heading {level - 1}')
    else:
        bullet = line.startswith('- ') and not code
        text = line[2:] if bullet else line
        p = doc.add_paragraph(text.replace('`', ''), style='List Bullet' if bullet else 'Normal')
        if code:
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
doc.save(target)
print('Word reference synchronized with reviewed Markdown')
